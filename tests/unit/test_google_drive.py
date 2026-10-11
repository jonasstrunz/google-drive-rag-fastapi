import io
from unittest.mock import MagicMock

from docx import Document
from pypdf import PdfWriter

from src.app.services.google_drive import GoogleDriveService


def test_list_files():
    mock_drive_resource = MagicMock()
    mock_files = [
        {
            "id": "file_1",
            "name": "Design Doc.gdoc",
            "mimeType": "application/vnd.google-apps.document",
            "size": "1024",
            "modifiedTime": "2026-10-11T12:00:00Z",
            "webViewLink": "https://docs.google.com/doc/1",
            "iconLink": "https://icon.png",
        },
        {
            "id": "file_2",
            "name": "Report.pdf",
            "mimeType": "application/pdf",
            "size": "2048",
            "modifiedTime": "2026-10-10T10:00:00Z",
            "webViewLink": "https://drive.google.com/file/2",
            "iconLink": "https://icon.png",
        },
    ]

    mock_drive_resource.files().list().execute.return_value = {"files": mock_files}

    service = GoogleDriveService(drive_resource=mock_drive_resource)
    files = service.list_files()

    assert len(files) == 2
    assert files[0].id == "file_1"
    assert files[0].name == "Design Doc.gdoc"
    assert files[1].size == 2048


def test_get_file_metadata():
    mock_drive_resource = MagicMock()
    mock_drive_resource.files().get().execute.return_value = {
        "id": "file_1",
        "name": "Readme.md",
        "mimeType": "text/markdown",
        "size": "500",
        "modifiedTime": "2026-10-11T14:00:00Z",
    }

    service = GoogleDriveService(drive_resource=mock_drive_resource)
    meta = service.get_file_metadata("file_1")

    assert meta.id == "file_1"
    assert meta.name == "Readme.md"
    assert meta.mime_type == "text/markdown"


def test_extract_google_doc_content(monkeypatch):
    mock_drive_resource = MagicMock()
    mock_drive_resource.files().get().execute.return_value = {
        "id": "doc_123",
        "name": "Strategy Plan",
        "mimeType": "application/vnd.google-apps.document",
    }

    # Mock export_media downloader
    export_content = b"This is the strategic roadmap for Q4."

    def fake_next_chunk():
        return None, True

    mock_downloader = MagicMock()
    mock_downloader.next_chunk.side_effect = fake_next_chunk

    def fake_export_downloader(fh, request):
        fh.write(export_content)
        return mock_downloader

    monkeypatch.setattr(
        "src.app.services.google_drive.MediaIoBaseDownload",
        fake_export_downloader,
    )

    service = GoogleDriveService(drive_resource=mock_drive_resource)
    text, meta = service.extract_file_content("doc_123")

    assert text == "This is the strategic roadmap for Q4."
    assert meta.id == "doc_123"


def test_extract_docx_content():
    # Generate real in-memory docx bytes
    doc = Document()
    doc.add_paragraph("First paragraph in DOCX.")
    doc.add_paragraph("Second paragraph in DOCX.")
    docx_io = io.BytesIO()
    doc.save(docx_io)
    docx_bytes = docx_io.getvalue()

    text = GoogleDriveService._extract_docx_text(docx_bytes)
    assert "First paragraph in DOCX." in text
    assert "Second paragraph in DOCX." in text


def test_extract_pdf_content():
    # Generate real in-memory PDF bytes
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    pdf_io = io.BytesIO()
    writer.write(pdf_io)
    pdf_bytes = pdf_io.getvalue()

    text = GoogleDriveService._extract_pdf_text(pdf_bytes)
    assert isinstance(text, str)
